"""Foglio "In breve" (0.4.0): primo foglio del registro condiviso, per i condomini da telefono."""
import datetime as dt
import os
import sys
import unittest

import openpyxl

from helpers import CondominioTest, RIPARTO, SKILLS, run

sys.path.insert(0, os.path.join(SKILLS, "condominio-base", "scripts"))
import registro  # noqa: E402

OGGI = dt.date.today()


def testo(ws):
    return "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)


class TestFormati(unittest.TestCase):
    def test_data_ed_euro_in_italiano(self):
        self.assertEqual(registro.data_it("2026-11-30"), "30 novembre 2026")
        self.assertEqual(registro.euro_it(1250.5), "1.250,50 €")
        self.assertEqual(registro.euro_it(-10), "-10,00 €")
        self.assertEqual(registro.euro_it(0.005), "0,01 €")

    def test_descrizione_leggibile_dal_nome_del_file(self):
        d = registro.descrivi_file
        self.assertEqual(d("2026-03-14_fattura_enel_luce-scale_412.50.pdf"),
                         (dt.date(2026, 3, 14), "Fattura Enel, luce scale · 412,50 €"))
        self.assertEqual(d("2026-01-10_contratto_otis_manutenzione-ascensore.pdf"),
                         (dt.date(2026, 1, 10), "Contratto Otis, manutenzione ascensore"))
        self.assertEqual(d("2026-02-02_preventivo_rossi-impianti_sostituzione-caldaia_8900.00.pdf")[1],
                         "Preventivo Rossi Impianti, sostituzione caldaia · 8.900,00 €")
        self.assertEqual(d("scansione senza nome.pdf"), (None, "scansione senza nome.pdf"))


class TestInBreve(CondominioTest):
    def in_breve(self):
        wb = self.wb()
        self.assertEqual(wb.sheetnames[0], "In breve")
        self.assertEqual(wb.active.title, "In breve")
        return wb["In breve"]

    def test_nuovo_registro_si_apre_su_in_breve(self):
        ws = self.in_breve()
        self.assertEqual(ws["A1"].value, "Condominio Prova")
        self.assertIn(registro.euro_it(2442.5), testo(ws))  # totale spese dell'esercizio di esempio

    def test_si_aggiorna_a_ogni_scrittura(self):
        self.reg("append", "Spese", '{"Data": "2026-06-01", "Fornitore": "X", "Importo": 57.5, "Tabella": "A", "Esercizio": 2026}')
        self.assertIn(registro.euro_it(2500), testo(self.in_breve()))
        self.assertEqual(self.wb().sheetnames.count("In breve"), 1)

    def test_prossime_scadenze_e_documenti_in_attesa(self):
        futura = OGGI + dt.timedelta(days=20)
        self.reg("append", "Scadenze", f'{{"Data": "{futura}", "Tipo": "assemblea", "Descrizione": "Assemblea ordinaria"}}')
        self.reg("append", "Scadenze", f'{{"Data": "{OGGI - dt.timedelta(days=1)}", "Tipo": "altro", "Descrizione": "Ieri"}}')
        inbox = self.path("da analizzare")
        for n in ("bolletta.pdf", "foto.jpg", "scorciatoia.gdoc"):
            open(os.path.join(inbox, n), "w").close()
        self.reg("verifica")
        t = testo(self.in_breve())
        self.assertIn(registro.data_it(futura), t)
        self.assertIn("Assemblea ordinaria", t)
        self.assertNotIn("Ieri", t)  # le scadenze passate non sono "prossime"
        self.assertEqual(registro.file_in_attesa(self.dir), 2)
        self.assertIn("2", [str(c.value) for row in self.in_breve().iter_rows() for c in row])

    def test_nessun_dato_per_unita(self):
        self.reg("append", "Spese", '{"Data": "2026-06-01", "Fornitore": "Fabbro", "Descrizione": "Serratura", "Importo": 80, "Tabella": "UNITA:U03", "Esercizio": 2026}')
        _, prev = self.riparto("--rate", 2, "--titolo", "Preventivo")
        futuro = OGGI + dt.timedelta(days=30)
        code, out = run(RIPARTO, "approva", "--dir", self.dir, "--prospetto", prev["prospetto"],
                        "--scadenze", f"{futuro},{futuro + dt.timedelta(days=90)}", "--approvato", "Michele")
        self.assertEqual(code, 0, out)
        t = testo(self.in_breve())
        _, anagrafica = self.reg("read", "Anagrafica")
        for u in anagrafica:
            for campo in ("ID unità", "Intestatario", "Email", "Conduttore", "Email conduttore"):
                if u.get(campo):
                    self.assertNotIn(str(u[campo]), t, f"{campo} nel foglio In breve")
        self.assertIn("Rata 1 di 2 — Preventivo", t)

    def test_sola_consultazione(self):
        for args in (("read", "In breve"), ("append", "In breve", "{}"), ("update", "In breve", "--where", "A=1", "{}")):
            code, out = self.reg(*args)
            self.assertNotEqual(code, 0, args)
        _, info = self.reg("info")
        self.assertNotIn("righe_In breve", info)

    def test_modifiche_a_mano_sovrascritte(self):
        wb = self.wb()
        wb["In breve"]["A1"] = "scritto a mano"
        wb.save(self.path("registro-condominio.xlsx"))
        self.reg("diario", "Prova")
        self.assertEqual(self.in_breve()["A1"].value, "Condominio Prova")

    def test_registro_0_3_aggiornato_con_schema(self):
        wb = self.wb()
        del wb["In breve"]
        wb.active = 0
        wb.save(self.path("registro-condominio.xlsx"))
        code, out = self.reg("schema")
        self.assertEqual(code, 0)
        self.assertEqual(out["aggiunte"], ["foglio 'In breve' (in prima posizione, generato)"])
        self.in_breve()
        self.assertEqual(self.reg("schema")[1]["aggiunte"], [])

    def test_un_solo_foglio_selezionato(self):
        """Più fogli selezionati insieme in Excel diventano un 'gruppo': le modifiche finirebbero su tutti."""
        self.reg("verifica")
        wb = self.wb()
        self.assertEqual([ws.title for ws in wb.worksheets if ws.sheet_view.tabSelected], ["In breve"])


if __name__ == "__main__":
    unittest.main()
