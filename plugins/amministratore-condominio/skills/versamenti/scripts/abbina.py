#!/usr/bin/env python3
"""
abbina.py — abbina i movimenti in entrata (bonifici) alle unità e alle rate, poi li registra.

  abbina.py --dir "<cartella>" --movimenti movimenti.json [--esercizio 2026] [--oggi 2026-09-26]
      Non scrive nulla. Legge un JSON [{"data", "importo", "ordinante", "causale"[, "unita"]}, ...]
      (estratto da un estratto conto o dettato dall'utente; "unita" solo se l'ha indicata l'utente) e
      propone per ogni movimento: esito (proposto | da_abbinare | gia_registrato | ignorato), unità,
      rata coperta, motivi del punteggio, candidati alternativi e la riga di Versamenti pronta.

  abbina.py registra --dir "<cartella>" --versamenti confermati.json --approvato "<nome>"
      Scrive in Versamenti le righe confermate (le "versamento" della proposta, eventualmente corrette),
      crea le righe mancanti di Situazione e annota nel Diario solo il numero dei versamenti.

  abbina.py ricorda --dir "<cartella>" --ordinante "<come nell'estratto>" --unita U03 --approvato "<nome>"
  abbina.py dimentica --dir "<cartella>" --ordinante "<…>" [--unita U03] --approvato "<nome>"
      Memoria di chi paga per chi (foglio Ordinanti del registro riservato): un ordinante ricordato
      pesa quanto un nome completo. Si ricorda solo dopo un sì dell'utente; dimenticare disattiva.

Il punteggio decide solo cosa PROPORRE: ogni abbinamento va confermato dall'utente.
"""
import argparse
import datetime as dt
import hashlib
import json
import os
import re
import sys
import unicodedata
from decimal import Decimal

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "condominio-base", "scripts"))
try:
    import openpyxl
    import registro as reg
except ImportError as e:  # pragma: no cover
    if e.name == "openpyxl":
        sys.exit("openpyxl non installato. Installarlo con: python3 -m pip install openpyxl (Windows: py -m pip install openpyxl; Linux con errore 'externally-managed-environment': sudo apt install python3-openpyxl)")
    sys.exit(f"registro.py non trovato ({e}): il plugin è installato in modo incompleto")

CENT = Decimal("0.01")

# Pesi e soglie. Principio: meglio un "da abbinare" in più che un bonifico sull'unità sbagliata.
# Un nome completo basta da solo; un cognome da solo no (omonimie, familiari); l'interno nella causale
# o l'importo esatto di una rata servono a confermare, non a decidere da soli.
PESO_NOME_COMPLETO = 60   # tutte le parole del nome (intestatario, comproprietario o conduttore) nel testo
PESO_NOME_PARZIALE = 25   # solo alcune (tipicamente il cognome)
PESO_INTERNO = 40         # "INT 3", "INTERNO 3", "APP 3" o l'ID unità nella causale
PESO_RATA = 30            # importo = parte ancora da pagare della prossima rata
PESO_SALDO = 20           # importo = tutto il saldo
PESO_ALTRA_RATA = 15      # importo = l'importo pieno di un'altra rata
PESO_MEMORIA = 60         # ordinante ricordato per l'unità (foglio Ordinanti): basta da solo, come un nome completo
SOGLIA = 60               # punteggio minimo per proporre
DISTACCO = 25             # vantaggio minimo sul secondo candidato


def norm(s):
    """Maiuscolo, senza accenti, solo lettere e cifre separate da uno spazio."""
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return " ".join(re.sub(r"[^A-Z0-9]+", " ", s.upper()).split())


def importo_da(v):
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v)).quantize(CENT)
    s = str(v or "").strip().replace("€", "").replace(" ", "")
    if "," in s:  # formato italiano 1.234,56
        s = s.replace(".", "").replace(",", ".")
    return Decimal(s).quantize(CENT)


def data_da(v):
    s = str(v or "").strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y", "%d.%m.%Y"):
        try:
            return dt.datetime.strptime(s[:10], fmt).date()
        except ValueError:
            pass
    raise ValueError(f"data non riconosciuta: {v!r} (atteso AAAA-MM-GG o GG/MM/AAAA)")


def impronte(movimenti):
    """ID movimento: data, importo e ordinante (non la causale, che il modello potrebbe trascrivere in modo
    diverso a una seconda lettura). Movimenti identici nello stesso elenco ricevono #2, #3…"""
    visti, out = {}, []
    for m in movimenti:
        base = f"{m['data'].isoformat()}|{m['importo']:.2f}|{norm(m['ordinante'])}"
        visti[base] = visti.get(base, 0) + 1
        h = hashlib.sha1(base.encode()).hexdigest()[:12]
        out.append(h if visti[base] == 1 else f"{h}#{visti[base]}")
    return out


def persone(anag_riga):
    """Nomi collegati a un'unità: intestatari (separati da ';') e conduttore."""
    nomi = str(anag_riga.get("Intestatario") or "").split(";") + [str(anag_riga.get("Conduttore") or "")]
    return [n.strip() for n in nomi if n.strip()]


def ordinanti_attivi(wb_r):
    """{ordinante normalizzato: {ID unità, …}} dalle righe attive del foglio Ordinanti."""
    memoria = {}
    if "Ordinanti" in wb_r.sheetnames:
        for r in reg._rows(wb_r["Ordinanti"]):
            if str(r.get("Attiva") or "SI").strip().upper() != "NO" and not reg._blank(r.get("Ordinante")):
                memoria.setdefault(norm(r["Ordinante"]), set()).add(str(r.get("ID unità") or "").strip())
    return memoria


def punteggio(mov, uid, anag_riga, prossima, saldo, importi_rate, ricordati=()):
    """(punti, motivi) di un movimento per un'unità. prossima = parte ancora da pagare della prima rata aperta;
    ricordati = unità per cui l'ordinante del movimento è nella memoria."""
    testo = f" {norm(mov['ordinante'])} {norm(mov['causale'])} "
    causale = f" {norm(mov['causale'])} "
    punti, motivi = 0, []
    if uid in ricordati:
        punti += PESO_MEMORIA
        motivi.append(f"ordinante ricordato per l'interno {anag_riga.get('Interno') or uid}")
    miglior_nome = 0
    for nome in persone(anag_riga):
        parole = [p for p in norm(nome).split() if len(p) >= 3]
        trovate = [p for p in parole if f" {p} " in testo]
        if parole and len(trovate) == len(parole):
            if PESO_NOME_COMPLETO > miglior_nome:
                miglior_nome, motivo = PESO_NOME_COMPLETO, f"nome completo ({nome})"
        elif trovate and PESO_NOME_PARZIALE > miglior_nome:
            miglior_nome, motivo = PESO_NOME_PARZIALE, f"solo in parte il nome ({' '.join(trovate).title()})"
    if miglior_nome:
        punti += miglior_nome
        motivi.append(motivo)
    interno = norm(anag_riga.get("Interno"))
    pattern = [rf" {re.escape(norm(uid))} "]
    if interno:
        pattern.append(rf" (INT|INTERNO|APP|APPT|APPARTAMENTO) ?0*{re.escape(interno)} ")
    if any(re.search(p, causale) for p in pattern):
        punti += PESO_INTERNO
        motivi.append(f"interno {anag_riga.get('Interno') or uid} nella causale")
    imp = mov["importo"]
    if prossima is not None and imp == prossima[1]:
        punti += PESO_RATA
        motivi.append(f"importo uguale alla rata {prossima[0]}")
    elif saldo is not None and saldo > 0 and imp == saldo:
        punti += PESO_SALDO
        motivi.append("importo uguale al saldo")
    elif imp in importi_rate:
        punti += PESO_ALTRA_RATA
        motivi.append("importo uguale a una rata")
    return punti, motivi


class Conti:
    """Rate aperte per unità, aggiornate man mano che i movimenti dello stesso elenco vengono imputati."""

    def __init__(self, sit):
        self.rate = {s["ID unità"]: [dict(r) for r in s["rate"]] for s in sit}
        self.saldo = {s["ID unità"]: Decimal(str(s["Saldo"])) for s in sit}
        self.esercizio = {s["ID unità"]: s["Esercizio"] for s in sit}

    def prossima(self, uid):
        for r in self.rate.get(uid, []):
            resto = Decimal(str(r["Importo"])) - Decimal(str(r["Pagato"]))
            if resto > 0:
                return r["Rata"], resto.quantize(CENT)
        return None

    def importi(self, uid):
        return {Decimal(str(r["Importo"])).quantize(CENT) for r in self.rate.get(uid, [])}

    def imputa(self, uid, importo):
        """Imputa l'importo alle rate aperte; restituisce (prima rata toccata, descrizione)."""
        prima, residuo, toccate = None, importo, []
        for r in self.rate.get(uid, []):
            resto = Decimal(str(r["Importo"])) - Decimal(str(r["Pagato"]))
            if resto <= 0 or residuo <= 0:
                continue
            quota = min(resto, residuo)
            r["Pagato"] = float(Decimal(str(r["Pagato"])) + quota)
            residuo -= quota
            prima = prima if prima is not None else r["Rata"]
            toccate.append(f"rata {r['Rata']}" + ("" if quota == resto else " (in parte)"))
        if uid in self.saldo:
            self.saldo[uid] -= importo
        if not toccate:
            return None, "nessuna rata aperta: resta come credito" if self.rate.get(uid) else "l'unità non ha rate registrate"
        return prima, ", ".join(toccate) + (f"; eccedenza {reg.euro_it(residuo)} a credito" if residuo > 0 else "")


def carica(a):
    path_c = os.path.join(a.dir, reg.DEFAULT_FILE)
    if not os.path.exists(path_c):
        sys.exit("registro-condominio.xlsx non trovato: eseguire condominio-setup")
    wb_c = openpyxl.load_workbook(path_c)
    path_r = reg.percorso_riservato(a.dir, wb_c)
    if not os.path.exists(path_r):
        sys.exit(f"Registro riservato non trovato: {path_r} (controllare 'Percorso registro riservato')")
    return wb_c, path_c, openpyxl.load_workbook(path_r), path_r


def proponi(a):
    wb_c, _, wb_r, _ = carica(a)
    esercizio = a.esercizio or int(reg._kv(wb_c["Condominio"]).get("Esercizio corrente") or a.oggi.year)
    anag = {str(r["ID unità"]).strip(): r for r in reg._rows(wb_c["Anagrafica"]) if not reg._blank(r.get("ID unità"))}
    conti = Conti(reg.situazione(wb_r, a.oggi, esercizio))
    gia = {str(v.get("ID movimento")).strip() for v in reg._rows(wb_r["Versamenti"]) if not reg._blank(v.get("ID movimento"))}
    memoria = ordinanti_attivi(wb_r)

    with open(a.movimenti, encoding="utf-8") as f:
        grezzi = json.load(f)
    movimenti, errori = [], []
    for i, m in enumerate(grezzi, start=1):
        try:
            movimenti.append({"n": i, "data": data_da(m.get("data")), "importo": importo_da(m.get("importo")),
                              "ordinante": str(m.get("ordinante") or "").strip(), "causale": str(m.get("causale") or "").strip(),
                              "unita": str(m.get("unita") or "").strip() or None})
        except (ValueError, ArithmeticError) as e:
            errori.append(f"Movimento {i}: {e}")
    if errori:
        print(json.dumps({"ok": False, "errori": errori}, ensure_ascii=False, indent=2))
        sys.exit(2)
    for m, h in zip(movimenti, impronte(movimenti)):
        m["id_movimento"] = h
    movimenti.sort(key=lambda m: (m["data"], m["n"]))  # imputazione in ordine cronologico

    out = []
    for m in movimenti:
        voce = {"n": m["n"], "data": m["data"].isoformat(), "importo": float(m["importo"]), "ordinante": m["ordinante"],
                "causale": m["causale"], "id_movimento": m["id_movimento"]}
        if m["importo"] <= 0:
            voce.update(esito="ignorato", motivo="addebito o importo nullo: la skill registra solo i versamenti in entrata")
        elif m["id_movimento"] in gia:
            voce.update(esito="gia_registrato", motivo="movimento già presente in Versamenti")
        elif m["unita"]:
            if m["unita"] not in anag:
                voce.update(esito="da_abbinare", motivo=f"unità {m['unita']} inesistente in Anagrafica", candidati=[])
            else:
                voce.update(esito="proposto", unita=m["unita"], intestatario=anag[m["unita"]].get("Intestatario"),
                            punteggio=None, motivi=["unità indicata dall'utente"])
        else:
            classifica = []
            for uid, riga in anag.items():
                p, motivi = punteggio(m, uid, riga, conti.prossima(uid), conti.saldo.get(uid), conti.importi(uid),
                                      memoria.get(norm(m["ordinante"]), ()))
                if p > 0:
                    classifica.append({"unita": uid, "intestatario": riga.get("Intestatario"), "punteggio": p, "motivi": motivi})
            classifica.sort(key=lambda c: -c["punteggio"])
            primo = classifica[0] if classifica else None
            secondo = classifica[1]["punteggio"] if len(classifica) > 1 else 0
            if primo and primo["punteggio"] >= SOGLIA and primo["punteggio"] - secondo >= DISTACCO:
                voce.update(esito="proposto", **primo, candidati=classifica[1:3])
            else:
                perche = ("nessun indizio" if not primo else
                          "indizi deboli" if primo["punteggio"] < SOGLIA else "più unità possibili")
                voce.update(esito="da_abbinare", motivo=perche, candidati=classifica[:3])
        if voce["esito"] == "proposto":
            uid = voce["unita"]
            rata, copre = conti.imputa(uid, m["importo"])
            voce.update(rata=rata, copre=copre, versamento={
                "Data": voce["data"], "ID unità": uid, "Importo": voce["importo"],
                "Esercizio": conti.esercizio.get(uid) or esercizio, "Rata": rata,
                "Riferimento": (m["causale"] or m["ordinante"])[:80], "ID movimento": m["id_movimento"]})
        out.append(voce)
    conteggio = {e: sum(v["esito"] == e for v in out) for e in ("proposto", "da_abbinare", "gia_registrato", "ignorato")}
    print(json.dumps({"ok": True, "esercizio": esercizio, "oggi": a.oggi.isoformat(), "riepilogo": conteggio,
                      "movimenti": out}, ensure_ascii=False, indent=2, default=str))


def registra(a):
    wb_c, path_c, wb_r, path_r = carica(a)
    with open(a.versamenti, encoding="utf-8") as f:
        righe = json.load(f)
    anag = {str(r["ID unità"]).strip(): r for r in reg._rows(wb_c["Anagrafica"]) if not reg._blank(r.get("ID unità"))}
    gia = {str(v.get("ID movimento")).strip() for v in reg._rows(wb_r["Versamenti"]) if not reg._blank(v.get("ID movimento"))}
    errori, nuove_ids = [], set()
    for i, v in enumerate(righe, start=1):
        uid = str(v.get("ID unità") or "").strip()
        if uid not in anag:
            errori.append(f"Versamento {i}: unità {uid or '(vuota)'} inesistente in Anagrafica")
        try:
            if importo_da(v.get("Importo")) <= 0:
                errori.append(f"Versamento {i}: importo non positivo")
            data_da(v.get("Data"))
        except (ValueError, ArithmeticError) as e:
            errori.append(f"Versamento {i}: {e}")
        mid = str(v.get("ID movimento") or "").strip()
        if mid and (mid in gia or mid in nuove_ids):
            errori.append(f"Versamento {i}: movimento {mid} già registrato")
        if mid:
            nuove_ids.add(mid)
    if errori:
        print(json.dumps({"ok": False, "errori": errori}, ensure_ascii=False, indent=2))
        sys.exit(2)

    ws_v, ws_s = wb_r["Versamenti"], wb_r["Situazione"]
    esistenti = {(str(r.get("ID unità") or "").strip(), int(reg._num(r.get("Esercizio")) or 0)) for r in reg._rows(ws_s)}
    ids, situazioni_create = [], []
    for v in righe:
        riga = {k: v.get(k) for k in ("Data", "ID unità", "Importo", "Esercizio", "Riferimento", "Rata", "Note", "ID movimento")
                if not reg._blank(v.get(k))}
        riga["Data"] = data_da(riga["Data"])
        riga["Importo"] = float(importo_da(riga["Importo"]))
        riga["Esercizio"] = int(riga.get("Esercizio") or riga["Data"].year)
        ids.append(reg.aggiungi(ws_v, riga)[1])
        chiave = (riga["ID unità"], riga["Esercizio"])
        if chiave not in esistenti:  # senza riga in Situazione il versamento non comparirebbe nel Versato
            reg.aggiungi(ws_s, {"ID unità": riga["ID unità"], "Intestatario": anag[riga["ID unità"]].get("Intestatario"),
                                "Esercizio": riga["Esercizio"], "Livello sollecito": 0})
            esistenti.add(chiave)
            situazioni_create.append(riga["ID unità"])
    # Diario pubblico: solo il numero, mai nomi o importi
    n = len(righe)
    reg.aggiungi(wb_c["Diario"], {"Data-ora": dt.datetime.now().isoformat(timespec="seconds"),
                                  "Operazione": f"Registrat{'o' if n == 1 else 'i'} {n} versament{'o' if n == 1 else 'i'}",
                                  "Dettaglio": "dettaglio nel registro riservato", "Eseguito da": "IA", "Approvato da": a.approvato})
    reg._save(wb_r, path_r, a.oggi)
    reg._save(wb_c, path_c, a.oggi)
    unita = sorted({str(v.get("ID unità")).strip() for v in righe})
    sit = [s for s in reg.situazione(openpyxl.load_workbook(path_r), a.oggi) if s["ID unità"] in unita]
    print(json.dumps({"ok": True, "registrati": n, "ID": ids, "situazioni_create": situazioni_create,
                      "situazione": [{k: s[k] for k in ("ID unità", "Esercizio", "Dovuto", "Versato", "Saldo", "Scaduto")} for s in sit]},
                     ensure_ascii=False, indent=2, default=str))


def _diario(wb_c, operazione, approvato):
    reg.aggiungi(wb_c["Diario"], {"Data-ora": dt.datetime.now().isoformat(timespec="seconds"), "Operazione": operazione,
                                  "Dettaglio": "dettaglio nel registro riservato", "Eseguito da": "IA", "Approvato da": approvato})


def ricorda(a):
    wb_c, path_c, wb_r, path_r = carica(a)
    if "Ordinanti" not in wb_r.sheetnames:
        sys.exit("Il registro riservato non ha il foglio Ordinanti: eseguire `registro.py --file registro-riservato.xlsx schema`")
    chiave, uid = norm(a.ordinante), a.unita.strip()
    anag = {str(r["ID unità"]).strip() for r in reg._rows(wb_c["Anagrafica"]) if not reg._blank(r.get("ID unità"))}
    errori = ([] if chiave else ["ordinante vuoto: si ricorda solo un ordinante scritto nell'estratto conto"]) + \
             ([] if uid in anag else [f"unità {uid} inesistente in Anagrafica"])
    if errori:
        print(json.dumps({"ok": False, "errori": errori}, ensure_ascii=False, indent=2))
        sys.exit(2)
    ws = wb_r["Ordinanti"]
    hdr = reg._headers(ws)
    c = {h: i for i, h in enumerate(hdr, start=1)}
    altre, riattivato = set(), False
    for r in range(2, ws.max_row + 1):
        if norm(ws.cell(r, c["Ordinante"]).value) != chiave:
            continue
        attiva = str(ws.cell(r, c["Attiva"]).value or "SI").strip().upper() != "NO"
        if str(ws.cell(r, c["ID unità"]).value or "").strip() == uid:
            if attiva:
                print(json.dumps({"ok": True, "ordinante": chiave, "unita": uid, "gia_ricordato": True}, ensure_ascii=False))
                return
            reg._put(ws, r, c["Attiva"], "Attiva", "SI")
            reg._put(ws, r, c["Dal"], "Dal", a.oggi)
            riattivato = True
        elif attiva:
            altre.add(str(ws.cell(r, c["ID unità"]).value or "").strip())
    if not riattivato:
        reg.aggiungi(ws, {"Ordinante": chiave, "ID unità": uid, "Dal": a.oggi, "Attiva": "SI"})
    _diario(wb_c, "Memorizzato 1 ordinante", a.approvato)
    reg._save(wb_r, path_r, a.oggi)
    reg._save(wb_c, path_c, a.oggi)
    out = {"ok": True, "ordinante": chiave, "unita": uid, "riattivato": riattivato, "anche_per": sorted(altre)}
    if altre:
        out["avviso"] = (f"{chiave} è ricordato anche per {', '.join(sorted(altre))}: i suoi bonifici resteranno "
                         "da abbinare se la causale non indica l'interno")
    print(json.dumps(out, ensure_ascii=False, indent=2))


def dimentica(a):
    wb_c, path_c, wb_r, path_r = carica(a)
    chiave = norm(a.ordinante)
    ws = wb_r["Ordinanti"] if "Ordinanti" in wb_r.sheetnames else None
    disattivate = []
    if ws is not None and chiave:
        hdr = reg._headers(ws)
        c = {h: i for i, h in enumerate(hdr, start=1)}
        for r in range(2, ws.max_row + 1):
            uid = str(ws.cell(r, c["ID unità"]).value or "").strip()
            if (norm(ws.cell(r, c["Ordinante"]).value) == chiave and (not a.unita or uid == a.unita.strip())
                    and str(ws.cell(r, c["Attiva"]).value or "SI").strip().upper() != "NO"):
                reg._put(ws, r, c["Attiva"], "Attiva", "NO")
                disattivate.append(uid)
    if not disattivate:
        print(json.dumps({"ok": False, "errori": [f"nessuna memoria attiva per {chiave or '(vuoto)'}"
                                                  + (f" e {a.unita}" if a.unita else "")]}, ensure_ascii=False, indent=2))
        sys.exit(2)
    n = len(disattivate)
    _diario(wb_c, f"Disattivat{'o' if n == 1 else 'i'} {n} ordinant{'e' if n == 1 else 'i'}", a.approvato)
    reg._save(wb_r, path_r, a.oggi)
    reg._save(wb_c, path_c, a.oggi)
    print(json.dumps({"ok": True, "ordinante": chiave, "disattivate": disattivate}, ensure_ascii=False, indent=2))


def main():
    comandi = {"registra": registra, "ricorda": ricorda, "dimentica": dimentica}
    cmd = sys.argv[1] if len(sys.argv) > 1 and sys.argv[1] in comandi else None
    p = argparse.ArgumentParser(prog="abbina.py" + (f" {cmd}" if cmd else ""), description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--dir", default=".")
    p.add_argument("--oggi", type=dt.date.fromisoformat, default=dt.date.today())
    if cmd == "registra":
        p.add_argument("--versamenti", required=True, help="JSON con le righe di Versamenti confermate")
    elif cmd in ("ricorda", "dimentica"):
        p.add_argument("--ordinante", required=True, help="ordinante come scritto nell'estratto conto")
        p.add_argument("--unita", required=cmd == "ricorda")
    else:
        p.add_argument("--movimenti", required=True, help="JSON con i movimenti da abbinare")
        p.add_argument("--esercizio", type=int)
    if cmd:
        p.add_argument("--approvato", required=True)
    a = p.parse_args(sys.argv[2:] if cmd else sys.argv[1:])
    try:
        (comandi[cmd] if cmd else proponi)(a)
    except SystemExit:
        raise
    except Exception as e:  # mai un traceback all'utente
        print(json.dumps({"ok": False, "errori": [f"{type(e).__name__}: {e}"]}, ensure_ascii=False, indent=2))
        sys.exit(2)


if __name__ == "__main__":
    main()
