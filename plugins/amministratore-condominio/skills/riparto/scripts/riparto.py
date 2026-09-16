#!/usr/bin/env python3
"""
riparto.py — calcola il riparto delle spese per unità secondo le tabelle millesimali.

  riparto.py --dir "<cartella>" [--esercizio 2026] [--ids 1,2,3] [--dal 2026-01-01] [--al 2026-06-30]
             [--titolo "Riparto 1° semestre 2026"] [--rate 4] [--out prospetti/riparto-2026-1sem.xlsx]

Regole:
- Tabella A/B/C/…: quota = importo × millesimi / 1000, arrotondata al centesimo.
- Il residuo di arrotondamento (max qualche centesimo) viene assegnato con il metodo del
  "resto maggiore" così che la somma delle quote sia ESATTAMENTE uguale all'importo.
- Importi negativi (note di credito, rimborsi) sono ripartiti con lo stesso criterio, con segno.
- UNITA:<ID>: tutto a quell'unità.
- MANUALE: quote lette dal foglio "Riparti manuali" (devono sommare all'importo).
- Si rifiuta di partire se una tabella usata non somma a 1000 (±0,01) o se una spesa non ha Tabella.

Produce un xlsx in prospetti/ (nome con suffisso _bozza; non sovrascrive mai un file esistente) con
fogli Riepilogo, Dettaglio, Millesimi usati — solo valori, nessuna formula — e stampa un riepilogo JSON.
"""
import argparse
import datetime as dt
import json
import os
import re
import sys
from decimal import Decimal, ROUND_DOWN

try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl non installato: pip install openpyxl")

FONT = "Arial"
CENT = Decimal("0.01")
FMT_EURO = "#,##0.00 €"
HDR_FILL = PatternFill("solid", fgColor="DDE7F0")


def D(x):
    if x is None or x == "":
        return Decimal(0)
    return Decimal(str(x).replace(",", "."))


def blank(v):
    return v is None or (isinstance(v, str) and not v.strip())


def headers(ws):
    out = []
    for c in ws[1]:
        if blank(c.value):
            break
        out.append(c.value)
    return out


def rows(ws):
    hdr = headers(ws)
    out = []
    for r in ws.iter_rows(min_row=2, max_col=len(hdr), values_only=True):
        if all(blank(v) for v in r):
            continue
        out.append(dict(zip(hdr, r)))
    return out


def as_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, str) and v:
        return dt.date.fromisoformat(v[:10])
    return None


def ripartisci(importo, pesi):
    """pesi: dict id->Decimal (millesimi o quote). Ritorna dict id->Decimal, somma esatta = importo."""
    importo = importo.quantize(CENT)
    if importo < 0:
        return {u: -q for u, q in ripartisci(-importo, pesi).items()}
    tot = sum(pesi.values())
    if tot == 0:
        raise ValueError("somma dei pesi nulla")
    grezze = {u: importo * p / tot for u, p in pesi.items()}
    troncate = {u: q.quantize(CENT, rounding=ROUND_DOWN) for u, q in grezze.items()}
    residuo = (importo - sum(troncate.values())).quantize(CENT)
    n_cent = int((residuo / CENT).to_integral_value())
    # assegna i centesimi mancanti alle unità con resto decimale maggiore (metodo di Hamilton)
    ordine = sorted(pesi, key=lambda u: (grezze[u] - troncate[u], pesi[u]), reverse=True)
    for i in range(n_cent):
        troncate[ordine[i % len(ordine)]] += CENT
    assert sum(troncate.values()) == importo, "riparto non quadra"
    return troncate


def slug(s, n=40):
    s = re.sub(r"[^a-z0-9]+", "-", s.lower().replace("à", "a").replace("è", "e").replace("é", "e").replace("ì", "i").replace("ò", "o").replace("ù", "u"))
    return s.strip("-")[:n].strip("-")


def nome_libero(path):
    """Non sovrascrive mai: aggiunge _2, _3… se il file esiste."""
    base, ext = os.path.splitext(path)
    k, cand = 1, path
    while os.path.exists(cand):
        k += 1
        cand = f"{base}_{k}{ext}"
    return cand


def fmt_row(ws, r, bold=False, euro_from=None):
    for c in range(1, ws.max_column + 1):
        cell = ws.cell(r, c)
        cell.font = Font(name=FONT, bold=bold)
        if euro_from and c >= euro_from and isinstance(cell.value, (int, float)):
            cell.number_format = FMT_EURO


def calcola(a):
    reg = os.path.join(a.dir, "registro-condominio.xlsx")
    if not os.path.exists(reg):
        sys.exit("registro-condominio.xlsx non trovato: eseguire condominio-setup")
    wb = openpyxl.load_workbook(reg)
    cond = {r[0].value: r[1].value for r in wb["Condominio"].iter_rows(min_row=1, max_col=2) if r[0].value}
    esercizio = a.esercizio or int(cond.get("Esercizio corrente") or dt.date.today().year)

    anag = {str(r["ID unità"]).strip(): r for r in rows(wb["Anagrafica"]) if not blank(r.get("ID unità"))}
    mill_rows = [r for r in rows(wb["Millesimi"]) if not blank(r.get("ID unità")) and str(r["ID unità"]).strip().upper() != "TOTALE"]
    tabelle = [k for k in headers(wb["Millesimi"]) if isinstance(k, str) and k.startswith("Tab ")]
    mill = {t: {str(r["ID unità"]).strip(): D(r.get(t)) for r in mill_rows} for t in tabelle}
    unita = [str(r["ID unità"]).strip() for r in mill_rows]
    if not unita:
        sys.exit("Nessuna unità nel foglio Millesimi: completare il setup")

    errori = []
    for t, pesi in mill.items():
        s = sum(pesi.values())
        if abs(s - 1000) > CENT and s != 0:
            errori.append(f"{t} somma a {s}, non a 1000")
    manuali = {}
    if "Riparti manuali" in wb.sheetnames:
        for r in rows(wb["Riparti manuali"]):
            if not blank(r.get("ID spesa")):
                manuali.setdefault(int(D(r["ID spesa"])), {})[str(r.get("ID unità")).strip()] = D(r.get("Quota"))

    # selezione spese
    spese = []
    ids = set(int(x) for x in a.ids.split(",")) if a.ids else None
    dal = dt.date.fromisoformat(a.dal) if a.dal else None
    al = dt.date.fromisoformat(a.al) if a.al else None
    for r in rows(wb["Spese"]):
        if blank(r.get("ID")) or blank(r.get("Importo")):
            continue
        if ids is not None and int(D(r["ID"])) not in ids:
            continue
        if ids is None and not blank(r.get("Esercizio")) and int(D(r["Esercizio"])) != esercizio:
            continue
        d = as_date(r.get("Data"))
        if dal and d and d < dal:
            continue
        if al and d and d > al:
            continue
        spese.append(r)
    if not spese:
        sys.exit("Nessuna spesa selezionata")

    # calcolo
    dettaglio = []  # (spesa, chiave tabella normalizzata, {unità: quota})
    for s in spese:
        importo = D(s["Importo"]).quantize(CENT)
        raw = str(s.get("Tabella") or "").strip()
        sid = int(D(s["ID"]))
        if not raw:
            errori.append(f"Spesa {sid} ({s.get('Descrizione')}): manca la Tabella (A, B, C…, UNITA:<ID> o MANUALE)"); continue
        up = raw.upper()
        if up.startswith("UNITA:") or up.startswith("UNITÀ:"):
            uid = raw.split(":", 1)[1].strip()
            if uid not in anag:
                errori.append(f"Spesa {sid}: unità {uid} inesistente"); continue
            chiave = f"UNITA:{uid}"
            quote = {u: (importo if u == uid else Decimal(0)) for u in unita}
        elif up == "MANUALE":
            q = manuali.get(sid)
            if not q:
                errori.append(f"Spesa {sid}: riparto MANUALE senza righe in 'Riparti manuali'"); continue
            if sum(q.values()).quantize(CENT) != importo:
                errori.append(f"Spesa {sid}: quote manuali sommano a {sum(q.values())}, importo {importo}"); continue
            chiave = "MANUALE"
            quote = {u: q.get(u, Decimal(0)).quantize(CENT) for u in unita}
        else:
            chiave = up
            key = f"Tab {up}"
            if key not in mill:
                errori.append(f"Spesa {sid}: tabella {raw} inesistente (disponibili: {', '.join(t[4:] for t in tabelle)})"); continue
            try:
                quote = ripartisci(importo, mill[key])
            except ValueError as e:
                errori.append(f"Spesa {sid}: {e}"); continue
        dettaglio.append((s, chiave, quote))
    if errori:
        print(json.dumps({"ok": False, "errori": errori}, ensure_ascii=False, indent=2))
        sys.exit(2)

    totali = {u: sum(q[u] for _, _, q in dettaglio) for u in unita}
    per_tab = {}
    for _, chiave, q in dettaglio:
        for u in unita:
            per_tab.setdefault(chiave, {}).setdefault(u, Decimal(0))
            per_tab[chiave][u] += q[u]
    per_tab = dict(sorted(per_tab.items()))
    totale = sum(D(s["Importo"]) for s, _, _ in dettaglio).quantize(CENT)
    assert sum(totali.values()) == totale
    rate = {u: ripartisci(totali[u], {i: Decimal(1) for i in range(a.rate)}) for u in unita} if a.rate else {}

    titolo = a.titolo or f"Riparto spese esercizio {esercizio}"
    oggi = dt.date.today().isoformat()
    nome = slug(titolo)
    nome = nome[len("riparto-"):] if nome.startswith("riparto-") else nome
    out = a.out or os.path.join("prospetti", f"{oggi}_riparto_{nome or f'esercizio-{esercizio}'}_bozza.xlsx")
    out_path = nome_libero(os.path.join(a.dir, out))
    out = os.path.relpath(out_path, a.dir)
    os.makedirs(os.path.dirname(out_path), exist_ok=True)

    # --- xlsx (solo valori: leggibile anche dall'anteprima di Google Drive su cellulare)
    xw = Workbook()
    ws = xw.active; ws.title = "Riepilogo"
    ws["A1"] = f"{cond.get('Nome', '')} — {titolo}"; ws["A1"].font = Font(name=FONT, bold=True, size=13)
    ws["A2"] = f"Generato il {oggi} · BOZZA: valida solo dopo approvazione dell'assemblea"; ws["A2"].font = Font(name=FONT, italic=True)
    hdr = ["ID unità", "Interno", "Intestatario"] + [f"Quota {t}" for t in per_tab] + ["Totale dovuto"] + [f"Rata {i + 1}" for i in range(a.rate)]
    ws.append([]); ws.append(hdr)
    for c in ws[4]:
        c.font = Font(name=FONT, bold=True); c.fill = HDR_FILL
    for u in unita:
        row = [u, anag.get(u, {}).get("Interno", ""), anag.get(u, {}).get("Intestatario", "")]
        row += [float(per_tab[t][u]) for t in per_tab] + [float(totali[u])]
        row += [float(rate[u][i]) for i in range(a.rate)]
        ws.append(row)
        fmt_row(ws, ws.max_row, euro_from=4)
    tot_row = ["TOTALE", "", ""] + [float(sum(v.values())) for v in per_tab.values()] + [float(totale)]
    tot_row += [float(sum(rate[u][i] for u in unita)) for i in range(a.rate)]
    ws.append(tot_row)
    fmt_row(ws, ws.max_row, bold=True, euro_from=4)
    for i, h in enumerate(hdr, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(h) + 2)
    ws.column_dimensions["C"].width = 30
    ws.freeze_panes = "D5"

    wd = xw.create_sheet("Dettaglio")
    wd.append(["ID spesa", "Data", "Fornitore", "Descrizione", "Importo", "Tabella"] + unita)
    for c in wd[1]:
        c.font = Font(name=FONT, bold=True); c.fill = HDR_FILL
    for s, chiave, q in dettaglio:
        wd.append([int(D(s["ID"])), as_date(s.get("Data")), s.get("Fornitore"), s.get("Descrizione"), float(D(s["Importo"])), chiave] + [float(q[u]) for u in unita])
        fmt_row(wd, wd.max_row, euro_from=5)
        wd.cell(wd.max_row, 2).number_format = "yyyy-mm-dd"
    wd.append(["TOTALE", "", "", "", float(totale), ""] + [float(totali[u]) for u in unita])
    fmt_row(wd, wd.max_row, bold=True, euro_from=5)
    wd.column_dimensions["D"].width = 36; wd.column_dimensions["C"].width = 22
    wd.freeze_panes = "G2"

    wm = xw.create_sheet("Millesimi usati")
    wm.append(["ID unità"] + tabelle)
    for c in wm[1]:
        c.font = Font(name=FONT, bold=True); c.fill = HDR_FILL
    for u in unita:
        wm.append([u] + [float(mill[t][u]) for t in tabelle]); fmt_row(wm, wm.max_row)
    wm.append(["TOTALE"] + [float(sum(mill[t].values())) for t in tabelle]); fmt_row(wm, wm.max_row, bold=True)
    xw.save(out_path)

    print(json.dumps({
        "ok": True, "titolo": titolo, "esercizio": esercizio, "prospetto": out,
        "spese": len(dettaglio), "totale": float(totale), "rate": a.rate,
        "per_unita": {u: {"interno": anag.get(u, {}).get("Interno"), "intestatario": anag.get(u, {}).get("Intestatario"),
                          "totale": float(totali[u]), "rate": [float(rate[u][i]) for i in range(a.rate)]} for u in unita},
        "per_tabella": {t: float(sum(v.values())) for t, v in per_tab.items()},
        "a_carico_singola_unita": {t: float(sum(v.values())) for t, v in per_tab.items() if t.startswith("UNITA:")},
    }, ensure_ascii=False, indent=2, default=str))


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--dir", default=".")
    p.add_argument("--esercizio", type=int)
    p.add_argument("--ids", help="ID spese separati da virgola")
    p.add_argument("--dal"); p.add_argument("--al")
    p.add_argument("--titolo")
    p.add_argument("--rate", type=int, default=0, help="numero di rate in cui suddividere il dovuto (0 = nessuna)")
    p.add_argument("--out")
    a = p.parse_args()
    try:
        calcola(a)
    except SystemExit:
        raise
    except Exception as e:  # mai un traceback all'utente: errore leggibile in JSON
        print(json.dumps({"ok": False, "errori": [f"{type(e).__name__}: {e}"]}, ensure_ascii=False, indent=2))
        sys.exit(2)


if __name__ == "__main__":
    main()
