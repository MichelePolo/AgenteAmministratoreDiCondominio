#!/usr/bin/env python3
"""
stato.py — cruscotto del condominio: come siamo messi e cosa fare adesso. Sola lettura.

  stato.py --dir "<cartella>" [--oggi 2026-09-26]

Raccoglie in un JSON: documenti in attesa in "da analizzare", problemi e avvisi del registro,
scadenze vicine e da riproporre, assemblea e termini di legge, prospetti in bozza, rate scadute
per unità (dal registro riservato), modalità collaudo; poi al massimo 3 azioni suggerite, in
ordine di urgenza. Non scrive nulla: né registri, né Diario.
"""
import argparse
import datetime as dt
import glob
import json
import os
import sys
from decimal import Decimal

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "condominio-base", "scripts"))
try:
    import openpyxl
    import registro as reg
except ImportError as e:  # pragma: no cover
    if e.name == "openpyxl":
        sys.exit("openpyxl non installato. Installarlo con: python3 -m pip install openpyxl (Windows: py -m pip install openpyxl; Linux con errore 'externally-managed-environment': sudo apt install python3-openpyxl)")
    sys.exit(f"registro.py non trovato ({e}): il plugin è installato in modo incompleto")

GIORNI_PROSSIME = 30
GIORNI_RENDICONTO = 180   # art. 1130 n. 10 c.c.: assemblea per il rendiconto entro 180 giorni dalla chiusura
GIORNI_CONVOCAZIONE = 5   # art. 66 disp. att. c.c.: avviso almeno 5 giorni prima della prima convocazione
GIORNI_INBOX_VECCHIA = 7


def inbox(cartella, oggi):
    base = os.path.join(cartella, reg.INBOX)
    date = []
    for radice, _, nomi in os.walk(base):
        for n in nomi:
            if not reg.da_ignorare(n):
                date.append(dt.date.fromtimestamp(os.path.getmtime(os.path.join(radice, n))))
    return {"in_attesa": len(date), "piu_vecchio_giorni": (oggi - min(date)).days if date else None}


def scadenze(wb, oggi):
    righe = []
    for r in reg._rows(wb["Scadenze"]) if "Scadenze" in wb.sheetnames else []:
        d = reg._as_date(r.get("Data"))
        if d:
            righe.append({"id": reg._norm(r.get("ID")), "data": d, "tipo": str(r.get("Tipo") or "").strip().lower(),
                          "descrizione": r.get("Descrizione") or "", "ricorrenza": str(r.get("Ricorrenza") or "nessuna").strip().lower(),
                          "evento": not reg._blank(r.get("ID evento"))})
    righe.sort(key=lambda x: x["data"])
    future = [x for x in righe if x["data"] >= oggi]
    # una ricorrente passata va riproposta alla data successiva, se non c'è già una riga uguale nel futuro
    descr_future = {(x["tipo"], x["descrizione"]) for x in future}
    da_riproporre = [x for x in righe if x["data"] < oggi and x["ricorrenza"] not in ("", "nessuna")
                     and (x["tipo"], x["descrizione"]) not in descr_future]
    return {
        "prossime": [{**x, "tra_giorni": (x["data"] - oggi).days} for x in future if (x["data"] - oggi).days <= GIORNI_PROSSIME],
        "da_riproporre": da_riproporre,
        "senza_evento": sum(1 for x in future if not x["evento"]),
        "assemblea": next((x for x in future if x["tipo"] == "assemblea"), None),
    }


def convocazione_inviata(wb, assemblea):
    """Cerca nel Diario una convocazione inviata nei 60 giorni prima dell'assemblea."""
    if not assemblea or "Diario" not in wb.sheetnames:
        return None
    dal = assemblea["data"] - dt.timedelta(days=60)
    for r in reg._rows(wb["Diario"]):
        quando = reg._as_date(r.get("Data-ora"))
        if quando and dal <= quando <= assemblea["data"] and "convocazion" in str(r.get("Operazione") or "").lower():
            return True
    return False


def bozze(cartella, oggi):
    out = []
    for f in sorted(glob.glob(os.path.join(cartella, "prospetti", "*_bozza*.xlsx"))):
        creato = reg._as_date(os.path.basename(f)[:10]) or dt.date.fromtimestamp(os.path.getmtime(f))
        out.append({"file": os.path.relpath(f, cartella), "giorni": (oggi - creato).days})
    return out


def pagamenti(cartella, wb_c, oggi, esercizio):
    path_r = reg.percorso_riservato(cartella, wb_c)
    if not os.path.exists(path_r):
        return None
    sit = reg.situazione(openpyxl.load_workbook(path_r, read_only=False), oggi, esercizio)
    ritardo = [s for s in sit if s["Scaduto"]]
    return {
        "unita_in_ritardo": len(ritardo),
        "totale_scaduto": float(sum(Decimal(str(s["Scaduto"])) for s in ritardo)),
        "dettaglio": [{"unita": s["ID unità"], "intestatario": s["Intestatario"], "scaduto": s["Scaduto"],
                       "livello_sollecito": s["Livello sollecito"], "ultimo_sollecito": s["Ultimo sollecito"]} for s in ritardo],
        "senza_rate": sum(1 for s in sit if s["Scaduto"] is None),
    }


def azioni_suggerite(st):
    """Al massimo 3 azioni, dalla più urgente. Criterio: prima ciò che blocca, poi i termini di legge,
    poi i soldi (registrando i bonifici prima di sollecitare), poi l'ordine della cartella.
    Ogni azione: priorita (1 = più urgente), testo, frase da dire all'assistente."""
    azioni = []

    def aggiungi(priorita, testo, frase):
        azioni.append({"priorita": priorita, "testo": testo, "frase": frase})

    if st["registro"]["problemi"]:
        n = len(st["registro"]["problemi"])
        aggiungi(1, f"Il registro ha {n} problem{'a' if n == 1 else 'i'} che bloccano riparti e verifiche: "
                    f"{st['registro']['problemi'][0]}", "controlla il registro")

    ass = st["scadenze"]["assemblea"]
    if ass and st["assemblea"]["convocazione_inviata"] is False:
        margine = (ass["data"] - st["oggi"]).days - GIORNI_CONVOCAZIONE
        if margine < 0:
            aggiungi(1, f"Assemblea il {reg.data_it(ass['data'])} e nessuna convocazione nel Diario: il termine di "
                        f"{GIORNI_CONVOCAZIONE} giorni è già passato, valutare un rinvio", "convoca l'assemblea")
        elif margine <= 10:
            aggiungi(2, f"Assemblea il {reg.data_it(ass['data'])}: la convocazione va inviata entro il "
                        f"{reg.data_it(ass['data'] - dt.timedelta(days=GIORNI_CONVOCAZIONE))}", "convoca l'assemblea")
    termine = st["assemblea"]["termine_rendiconto"]
    if termine and not ass:
        mancano = (termine - st["oggi"]).days
        aggiungi(2 if mancano <= 60 else 4,
                 f"Esercizio chiuso: l'assemblea per il rendiconto va tenuta entro il {reg.data_it(termine)}"
                 + (" (termine superato)" if mancano < 0 else f" (mancano {mancano} giorni)"), "prepara il rendiconto")

    pag = st["pagamenti"]
    if pag and pag["unita_in_ritardo"]:
        n = pag["unita_in_ritardo"]
        aggiungi(3, f"{n} unit{'à ha' if n == 1 else 'à hanno'} rate scadute per {reg.euro_it(pag['totale_scaduto'])}: "
                    "prima registrare i bonifici arrivati, poi valutare i solleciti", "registra i bonifici dell'estratto conto")

    ib = st["inbox"]
    if ib["in_attesa"]:
        vecchio = (ib["piu_vecchio_giorni"] or 0) > GIORNI_INBOX_VECCHIA
        aggiungi(3 if vecchio else 4, f"{ib['in_attesa']} document{'o' if ib['in_attesa'] == 1 else 'i'} da archiviare"
                 + (f", il più vecchio caricato {ib['piu_vecchio_giorni']} giorni fa" if vecchio else ""), "archivia i documenti")

    if st["scadenze"]["da_riproporre"]:
        n = len(st["scadenze"]["da_riproporre"])
        aggiungi(4, f"{n} scadenz{'a ricorrente passata' if n == 1 else 'e ricorrenti passate'} da riproporre", "aggiorna le scadenze")
    elif st["scadenze"]["senza_evento"]:
        aggiungi(5, f"{st['scadenze']['senza_evento']} scadenze non ancora in calendario", "metti in calendario le scadenze")

    vecchie = [b for b in st["bozze"] if b["giorni"] > 60]
    if vecchie:
        aggiungi(5, f"{len(vecchie)} prospett{'o' if len(vecchie) == 1 else 'i'} in bozza da più di due mesi: "
                    "approvati in assemblea?", "l'assemblea ha approvato il riparto")

    azioni.sort(key=lambda a: a["priorita"])  # ordinamento stabile: a parità, l'ordine sopra
    return azioni[:3]


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--dir", default=".")
    p.add_argument("--oggi", type=dt.date.fromisoformat, default=dt.date.today())
    a = p.parse_args()
    path_c = os.path.join(a.dir, reg.DEFAULT_FILE)
    if not os.path.exists(path_c):
        sys.exit("registro-condominio.xlsx non trovato: la cartella di lavoro non è quella del condominio, o manca il setup")
    wb_c = openpyxl.load_workbook(path_c)
    kv = reg._kv(wb_c["Condominio"])
    esercizio = int(reg._num(kv.get("Esercizio corrente")) or a.oggi.year)
    fine = reg._as_date(kv.get("Fine esercizio"))
    problemi, avvisi = reg.verifica(wb_c)
    sc = scadenze(wb_c, a.oggi)
    st = {
        "oggi": a.oggi,
        "condominio": {"nome": kv.get("Nome"), "profilo": kv.get("Profilo"), "esercizio": esercizio,
                       "collaudo": str(kv.get("Modalità collaudo") or "").strip().upper() == "SI",
                       "ultima_elaborazione": kv.get("Ultima elaborazione")},
        "inbox": inbox(a.dir, a.oggi),
        "registro": {"problemi": problemi, "avvisi": avvisi},
        "scadenze": sc,
        "assemblea": {"convocazione_inviata": convocazione_inviata(wb_c, sc["assemblea"]),
                      "termine_rendiconto": fine + dt.timedelta(days=GIORNI_RENDICONTO) if fine and fine < a.oggi else None},
        "bozze": bozze(a.dir, a.oggi),
        "pagamenti": pagamenti(a.dir, wb_c, a.oggi, esercizio),
    }
    if st["pagamenti"] is None:
        avvisi.append("Registro riservato non trovato: pagamenti non verificati (controllare 'Percorso registro riservato')")
    st["azioni"] = azioni_suggerite(st)
    print(json.dumps(st, ensure_ascii=False, indent=2, default=lambda o: o.isoformat() if hasattr(o, "isoformat") else str(o)))


if __name__ == "__main__":
    main()
