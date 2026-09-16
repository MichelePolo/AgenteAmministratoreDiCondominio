#!/usr/bin/env python3
"""
scansiona.py — elenca i file in "da analizzare/", ne calcola l'hash e li confronta con il foglio Indice.

  scansiona.py --dir "<cartella condominio>"

Output JSON: per ogni file nome, dimensione, estensione, hash, stato (nuovo | duplicato | duplicato_in_inbox)
e, se duplicato, dove si trova la copia già archiviata. Non sposta nulla.
"""
import argparse
import hashlib
import json
import os
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl non installato: pip install openpyxl")

INBOX = "da analizzare"
SKIP_EXT = {".gdoc", ".gsheet", ".gslides", ".gform", ".gdraw", ".gmap", ".tmp", ".crdownload", ".drivedownload", ".driveupload"}


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def indice(reg_path):
    if not os.path.exists(reg_path):
        return {}
    wb = openpyxl.load_workbook(reg_path, read_only=True)
    if "Indice" not in wb.sheetnames:
        return {}
    ws = wb["Indice"]
    hdr = []
    for v in next(ws.iter_rows(min_row=1, max_row=1, values_only=True)):
        if v is None or (isinstance(v, str) and not v.strip()):
            break
        hdr.append(v)
    out = {}
    for r in ws.iter_rows(min_row=2, max_col=len(hdr), values_only=True):
        row = dict(zip(hdr, r))
        if row.get("Hash"):
            out[row["Hash"]] = row
    return out


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--dir", default=".")
    a = p.parse_args()
    inbox = os.path.join(a.dir, INBOX)
    if not os.path.isdir(inbox):
        sys.exit(f"Cartella '{INBOX}' non trovata in {a.dir}. Eseguire condominio-setup.")
    idx = indice(os.path.join(a.dir, "registro-condominio.xlsx"))
    files = []
    visti = {}  # hash -> primo file della stessa scansione
    for root, _, names in os.walk(inbox):
        for n in sorted(names):
            if n.startswith(".") or n.startswith("~$") or n.lower() in ("desktop.ini", "thumbs.db"):
                continue
            if os.path.splitext(n)[1].lower() in SKIP_EXT:  # scorciatoie e file temporanei di Google Drive
                continue
            fp = os.path.join(root, n)
            h = sha256(fp)
            rel = os.path.relpath(fp, a.dir)
            stato = "nuovo"
            copia = None
            if h in idx:
                stato = "duplicato"
                copia = idx[h].get("Percorso")
            elif h in visti:
                stato = "duplicato_in_inbox"
                copia = visti[h]
            else:
                visti[h] = rel
            files.append({
                "nome": n, "percorso": rel, "dimensione": os.path.getsize(fp),
                "estensione": os.path.splitext(n)[1].lower().lstrip("."),
                "hash": h, "stato": stato, "copia_archiviata": copia,
            })
    print(json.dumps({"inbox": INBOX, "totale": len(files), "nuovi": sum(f["stato"] == "nuovo" for f in files), "file": files}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
