#!/usr/bin/env python3
"""
crea_registro.py — crea la struttura della cartella condominio e i due registri xlsx.

  crea_registro.py --dir "<cartella condominio>" [--nome "Condominio X"] [--esercizio 2026] [--esempio]

--esempio inserisce dati dimostrativi (4 unità di cui una affittata, 3 spese, 2 versamenti) per provare subito un riparto.
Senza --esempio, Anagrafica e Millesimi nascono vuoti: le unità vere si inseriscono con
registro.py append (o a mano nel foglio). Non sovrascrive registri esistenti.

I registri non contengono formule: la riga TOTALE dei millesimi e le colonne Versato/Saldo sono
numeri, ricalcolati da registro.py a ogni operazione. Così anche l'anteprima di Google Drive su
cellulare mostra i valori giusti.
"""
import argparse
import datetime as dt
import os
import shutil
import sys

try:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl non installato: pip install openpyxl")

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="DDE7F0")
INPUT_FILL = PatternFill("solid", fgColor="FFF8DC")
FMT_EURO = "#,##0.00 €"
FMT_DATE = "yyyy-mm-dd"
COL_IMPORTI = {"Importo", "Quota", "Dovuto", "Versato", "Saldo"}
COL_DATE = {"Data", "Data pagamento", "Ultimo sollecito", "Data elaborazione"}

FOLDERS = ["da analizzare", "prospetti"] + [
    f"archivio/{{anno}}/{s}" for s in ("fatture", "preventivi", "verbali", "comunicazioni", "contratti", "altro", "altro/duplicati")
]


def _sheet(wb, name, headers, widths=None):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True)
        c.fill = HDR_FILL
        c.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, max(14, len(h) + 2))
    return ws


def _row(ws, headers, values, bold=False):
    """Aggiunge una riga applicando font e formati numerici/data in base all'intestazione."""
    ws.append(values)
    r = ws.max_row
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(r, c)
        cell.font = Font(name=FONT, bold=bold)
        if h in COL_IMPORTI:
            cell.number_format = FMT_EURO
        elif h in COL_DATE:
            cell.number_format = FMT_DATE


def crea_condominio(path, nome, esercizio, esempio):
    wb = Workbook()
    ws = wb.active
    ws.title = "Condominio"
    kv = [
        ("Nome", nome, "es. Condominio Le Betulle"),
        ("Indirizzo", "", "via e città"),
        ("Codice fiscale", "", "del condominio; obbligatorio se c'è un amministratore"),
        ("Profilo", "autogestione", "autogestione oppure professionista"),
        ("Amministratore", "", "nome e cognome"),
        ("Email amministratore", "", "mittente delle comunicazioni"),
        ("Modalità collaudo", "SI", "SI = ogni email va solo a 'Email collaudo'"),
        ("Email collaudo", "", "di solito quella dell'amministratore"),
        ("Esercizio corrente", esercizio, "anno di gestione in corso"),
        ("Inizio esercizio", f"{esercizio}-01-01", "data ISO"),
        ("Fine esercizio", f"{esercizio}-12-31", "data ISO"),
        ("Numero unità", 4 if esempio else "", "coerente con il foglio Anagrafica"),
        ("Tabelle millesimali", "A: proprietà generale; B: scale; C: ascensore", "cosa rappresentano le colonne Tab di Millesimi"),
        ("Conto corrente", "", "IBAN intestato al condominio (art. 1129 c.c.)"),
        ("Percorso registro riservato", "", "vuoto = stessa cartella; altrimenti la cartella privata dove sta registro-riservato.xlsx"),
        ("Ultima elaborazione", dt.datetime.now().isoformat(timespec="seconds"), "aggiornata dagli script"),
    ]
    for k, v, nota in kv:
        ws.append([k, v, nota])
    for r in ws.iter_rows(min_row=1, max_col=3):
        r[0].font = Font(name=FONT, bold=True); r[0].fill = HDR_FILL
        r[1].font = Font(name=FONT); r[1].fill = INPUT_FILL
        r[2].font = Font(name=FONT, italic=True, color="666666")
    ws.column_dimensions["A"].width = 26; ws.column_dimensions["B"].width = 42; ws.column_dimensions["C"].width = 60
    ws.freeze_panes = "B1"

    an_h = ["ID unità", "Interno", "Piano", "Intestatario", "Email", "Telefono", "Conduttore", "Email conduttore", "Dati catastali", "Note"]
    an = _sheet(wb, "Anagrafica", an_h, {"Intestatario": 28, "Email": 30, "Conduttore": 24, "Email conduttore": 30, "Dati catastali": 24, "Note": 30})
    mi_h = ["ID unità", "Intestatario", "Tab A", "Tab B", "Tab C"]
    mi = _sheet(wb, "Millesimi", mi_h, {"Intestatario": 28})
    sp_h = ["ID", "Data", "Fornitore", "Descrizione", "Importo", "Tabella", "Quota conduttore %", "Esercizio", "File", "Pagata", "Data pagamento", "Note"]
    sp = _sheet(wb, "Spese", sp_h, {"Descrizione": 34, "File": 60, "Fornitore": 22, "Quota conduttore %": 20, "Note": 30})
    sp["G1"].comment = Comment("Percentuale della quota di ogni unità che, all'interno dell'unità, spetta al conduttore "
                               "(art. 9 L. 392/1978): 100 per pulizie, luce scale, acqua, riscaldamento, ordinaria ascensore; "
                               "90 portierato; 0 per straordinaria, amministratore, assicurazione. Vuoto = 0. "
                               "Conta solo per le unità con Conduttore in Anagrafica.", "amministratore-condominio")
    _sheet(wb, "Riparti manuali", ["ID spesa", "ID unità", "Quota", "Note"])
    _sheet(wb, "Scadenze", ["ID", "Data", "Ora", "Tipo", "Descrizione", "Ricorrenza", "ID evento", "Note"], {"Descrizione": 40, "ID evento": 30, "Note": 30})
    di_h = ["Data-ora", "Operazione", "Dettaglio", "Eseguito da", "Approvato da"]
    di = _sheet(wb, "Diario", di_h, {"Data-ora": 20, "Operazione": 36, "Dettaglio": 70})
    _sheet(wb, "Indice", ["Hash", "Nome originale", "Nome archivio", "Percorso", "Data elaborazione", "ID spesa"], {"Hash": 20, "Nome originale": 40, "Nome archivio": 50, "Percorso": 50})

    unita = [("U01", "1", 0, "Mario Rossi", "mario.rossi@example.it", 300, 100, 0, "", ""),
             ("U02", "2", 1, "Anna Bianchi", "anna.bianchi@example.it", 250, 250, 300, "Paolo Gialli", "paolo.gialli@example.it"),
             ("U03", "3", 2, "Luca Verdi", "luca.verdi@example.it", 250, 300, 350, "", ""),
             ("U04", "4", 2, "Giulia Neri", "giulia.neri@example.it", 200, 350, 350, "", "")] if esempio else []
    for u in unita:
        _row(an, an_h, [u[0], u[1], u[2], u[3], u[4], "", u[8], u[9], "", ""])
        _row(mi, mi_h, [u[0], u[3], u[5], u[6], u[7]])
    tot = [sum(u[i] for u in unita) for i in (5, 6, 7)]
    _row(mi, mi_h, ["TOTALE", "ogni colonna Tab deve fare 1000", *tot], bold=True)
    mi["A1"].comment = Comment("Una riga per unità, sopra la riga TOTALE. Ogni colonna Tab deve sommare a 1000. "
                               "Altre tabelle (Tab D…) si aggiungono in coda con la stessa regola. "
                               "La riga TOTALE è ricalcolata dagli script.", "amministratore-condominio")
    if esempio:
        _row(sp, sp_h, [1, dt.date(esercizio, 3, 14), "ENEL", "Energia elettrica scale (esempio)", 412.50, "B", 100, esercizio, f"archivio/{esercizio}/fatture/{esercizio}-03-14_fattura_enel_luce-scale_412.50.pdf", "NO", None, "riga di esempio"])
        _row(sp, sp_h, [2, dt.date(esercizio, 2, 2), "Otis", "Manutenzione ordinaria ascensore 1° sem. (esempio)", 780.00, "C", 100, esercizio, f"archivio/{esercizio}/contratti/{esercizio}-01-10_contratto_otis_manutenzione-ascensore.pdf", "SI", dt.date(esercizio, 2, 20), "riga di esempio"])
        _row(sp, sp_h, [3, dt.date(esercizio, 4, 5), "Studio Tecnico Bruni", "Polizza globale fabbricato (esempio)", 1250.00, "A", 0, esercizio, "", "NO", None, "riga di esempio"])
    _row(di, di_h, [dt.datetime.now().isoformat(timespec="seconds"), "Creazione registro",
                    f"Registro creato da condominio-setup{' con dati di esempio' if esempio else ''}", "IA", ""])
    wb.save(path)


def crea_riservato(path, esercizio, esempio):
    wb = Workbook()
    ws = wb.active
    ws.title = "Situazione"
    hdr = ["ID unità", "Intestatario", "Esercizio", "Dovuto", "Versato", "Saldo", "Ultimo sollecito", "Livello sollecito", "Note"]
    ws.append(hdr)
    for c in ws[1]:
        c.font = Font(name=FONT, bold=True); c.fill = HDR_FILL
    ws.freeze_panes = "A2"
    ws["A1"].comment = Comment("Dovuto: aggiornato dopo l'approvazione di un prospetto. Versato e Saldo: ricalcolati da "
                               "registro.py dai Versamenti (Saldo > 0 = deve ancora versare). "
                               "Livello sollecito: 0 nessuno, 1 cortese, 2 formale, 3 diffida.", "amministratore-condominio")
    ve_h = ["ID", "Data", "ID unità", "Importo", "Esercizio", "Riferimento", "Rata", "Note"]
    ve = _sheet(wb, "Versamenti", ve_h, {"Riferimento": 30, "Note": 30})
    _sheet(wb, "Solleciti", ["ID", "Data", "ID unità", "Livello", "Inviato a", "Canale", "Esito", "Note"], {"Inviato a": 30, "Note": 30})
    unita = [("U01", "Mario Rossi"), ("U02", "Anna Bianchi"), ("U03", "Luca Verdi"), ("U04", "Giulia Neri")] if esempio else []
    versamenti = [(1, dt.date(esercizio, 1, 15), "U01", 150.00, esercizio, "bonifico rata 1", 1, "esempio"),
                  (2, dt.date(esercizio, 1, 16), "U02", 150.00, esercizio, "bonifico rata 1", 1, "esempio")] if esempio else []
    for uid, nome in unita:
        versato = round(sum(v[3] for v in versamenti if v[2] == uid), 2)
        _row(ws, hdr, [uid, nome, esercizio, 0, versato, round(0 - versato, 2), None, 0, ""])
    for v in versamenti:
        _row(ve, ve_h, list(v))
    for i, h in enumerate(hdr, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(h) + 4)
    wb.save(path)


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--dir", required=True)
    p.add_argument("--nome", default="Condominio")
    p.add_argument("--esercizio", type=int, default=dt.date.today().year)
    p.add_argument("--esempio", action="store_true")
    a = p.parse_args()
    os.makedirs(a.dir, exist_ok=True)
    for f in FOLDERS:
        os.makedirs(os.path.join(a.dir, f.format(anno=a.esercizio)), exist_ok=True)
    creati = []
    rc = os.path.join(a.dir, "registro-condominio.xlsx")
    rr = os.path.join(a.dir, "registro-riservato.xlsx")
    if not os.path.exists(rc):
        crea_condominio(rc, a.nome, a.esercizio, a.esempio); creati.append(rc)
    if not os.path.exists(rr):
        crea_riservato(rr, a.esercizio, a.esempio); creati.append(rr)
    leggimi_src = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "assets", "LEGGIMI.txt")
    leggimi_dst = os.path.join(a.dir, "LEGGIMI.txt")
    if os.path.exists(leggimi_src) and not os.path.exists(leggimi_dst):
        shutil.copy(leggimi_src, leggimi_dst); creati.append(leggimi_dst)
    print("Creati:" if creati else "Nulla da creare (registri già presenti).")
    for c in creati:
        print("  ", c)
    print("Ricordare: NON condividere registro-riservato.xlsx con i condomini.")


if __name__ == "__main__":
    main()
